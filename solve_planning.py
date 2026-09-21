"""Optimisation du planning hebdomadaire du service de pediatrie."""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import openpyxl
import pulp


ROOT = Path(__file__).resolve().parent
INPUT_FILE = ROOT / "Copie de Projet Plannification Hopital.xlsx"
OUTPUT_FILE = ROOT / "Planning_optimise.xlsx"

DAYS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
SLOTS = ["AM", "PM", "GA"]
ACTIVITY_COLUMNS = {
    "Pédiatrie": "Pédiatrie",
    "Urgences": "Urgences",
    "Consultations": "Consultations",
    "Exploration": "Exploration",
    "Astreinte": "Astreinte",
    "Maternité": "Maternité",
    "Néonatologie": "Néonatologie",
}
PRIORITY_ACTIVITIES = {
    "Pédiatrie",
    "Urgences",
    "Astreinte",
    "Maternité",
    "Néonatologie",
}


def normalize(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text.strip().lower())


def parse_demand(value: object) -> int:
    if value is None or value == "":
        return 0
    return int(value)


def read_data(path: Path) -> tuple[list[dict], dict, dict]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    staffing_sheet = workbook["2. Effectifs"]
    demand_sheet = workbook["1. Besoins"]
    absence_sheet = workbook["3. Contraintes"]

    doctors = []
    for row in staffing_sheet.iter_rows(min_row=5, max_col=12, values_only=True):
        if row[1] is None:
            continue
        name = str(row[4]).strip()
        skills = {
            activity
            for index, activity in enumerate(ACTIVITY_COLUMNS.values(), start=5)
            if index < len(row) and normalize(row[index]) == "x"
        }
        doctors.append({"name": name, "skills": skills})

    demand = {}
    for row in demand_sheet.iter_rows(min_row=7, max_row=13, min_col=2, max_col=23, values_only=True):
        activity = str(row[0]).strip()
        for slot_index, value in enumerate(row[1:]):
            day = DAYS[slot_index // 3]
            slot = SLOTS[slot_index % 3]
            demand[activity, day, slot] = parse_demand(value)

    absence_names = {}
    for row in absence_sheet.iter_rows(min_row=7, max_col=23, values_only=True):
        if row[1] is None or not str(row[1]).strip():
            continue
        name = normalize(row[1])
        for slot_index, value in enumerate(row[2:]):
            if slot_index >= len(DAYS) * len(SLOTS):
                break
            day = DAYS[slot_index // 3]
            slot = SLOTS[slot_index % 3]
            absence_names[name, day, slot] = normalize(value) == "vac"

    # Les noms sont parfois ecrits dans un ordre different entre les feuilles.
    absences = {}
    for doctor in doctors:
        doctor_name = normalize(doctor["name"])
        parts = doctor_name.split()
        reversed_name = " ".join(reversed(parts))
        for day in DAYS:
            for slot in SLOTS:
                absences[doctor["name"], day, slot] = (
                    absence_names.get((doctor_name, day, slot), False)
                    or absence_names.get((reversed_name, day, slot), False)
                )
    return doctors, demand, absences


def build_model(doctors: list[dict], demand: dict, absences: dict):
    model = pulp.LpProblem("Planning_Pediatrie", pulp.LpMinimize)
    doctor_names = [doctor["name"] for doctor in doctors]

    assignment = {}
    replacement = {}
    for activity in ACTIVITY_COLUMNS:
        for day in DAYS:
            for slot in SLOTS:
                if demand.get((activity, day, slot), 0) == 0:
                    continue
                replacement[activity, day, slot] = pulp.LpVariable(
                    f"remplacant_{activity}_{day}_{slot}", lowBound=0, cat="Integer"
                )
                for doctor in doctors:
                    name = doctor["name"]
                    if activity not in doctor["skills"] or absences[name, day, slot]:
                        continue
                    assignment[name, activity, day, slot] = pulp.LpVariable(
                        f"affectation_{len(assignment)}", cat="Binary"
                    )

    workload = {
        name: pulp.lpSum(
            variable
            for (doctor, _activity, _day, _slot), variable in assignment.items()
            if doctor == name
        )
        for name in doctor_names
    }
    average_workload = pulp.lpSum(workload.values()) / len(doctor_names)
    workload_deviation = {
        name: pulp.LpVariable(f"ecart_charge_{name}", lowBound=0)
        for name in doctor_names
    }

    for name in doctor_names:
        model += workload[name] <= 10, f"maximum_hebdomadaire_{name}"
        model += workload[name] - average_workload <= workload_deviation[name]
        model += average_workload - workload[name] <= workload_deviation[name]

    for activity in ACTIVITY_COLUMNS:
        for day in DAYS:
            for slot in SLOTS:
                required = demand.get((activity, day, slot), 0)
                if required == 0:
                    continue
                assigned = pulp.lpSum(
                    variable
                    for (doctor, current_activity, current_day, current_slot), variable in assignment.items()
                    if current_activity == activity and current_day == day and current_slot == slot
                )
                model += (
                    assigned + replacement[activity, day, slot] == required,
                    f"couverture_{activity}_{day}_{slot}",
                )

    for name in doctor_names:
        for day in DAYS:
            model += (
                pulp.lpSum(
                    variable
                    for (doctor, _activity, current_day, _slot), variable in assignment.items()
                    if doctor == name and current_day == day
                )
                <= 1,
                f"un_creneau_par_jour_{name}_{day}",
            )

    # Une garde impose le repos sur tous les creneaux du jour suivant.
    for name in doctor_names:
        for day_index, day in enumerate(DAYS):
            next_day = DAYS[(day_index + 1) % len(DAYS)]
            night_work = pulp.lpSum(
                variable
                for (doctor, _activity, current_day, current_slot), variable in assignment.items()
                if doctor == name and current_day == day and current_slot == "GA"
            )
            next_day_work = pulp.lpSum(
                variable
                for (doctor, _activity, current_day, _slot), variable in assignment.items()
                if doctor == name and current_day == next_day
            )
            model += night_work + next_day_work <= 1, f"repos_apres_garde_{name}_{day}"

    # Le premier terme rend le nombre de remplacants prioritaire. Le second
    # favorise les postes prioritaires a nombre de remplacants egal.
    replacement_cost = pulp.lpSum(
        variable * (1000 if activity in PRIORITY_ACTIVITIES else 100)
        for (activity, _day, _slot), variable in replacement.items()
    )
    model += (
        100000 * pulp.lpSum(replacement.values())
        + replacement_cost
        + 10 * pulp.lpSum(workload_deviation.values())
    )
    return model, assignment, replacement, workload


def export_solution(path: Path, model, assignment, replacement, workload):
    workbook = openpyxl.Workbook()
    schedule = workbook.active
    schedule.title = "Planning"
    schedule.append(["Jour", "Créneau", "Activité", "Médecin", "Type"])
    for (doctor, activity, day, slot), variable in assignment.items():
        if variable.value() == 1:
            schedule.append([day, slot, activity, doctor, "Titulaire"])
    for (activity, day, slot), variable in replacement.items():
        for _ in range(round(variable.value() or 0)):
            schedule.append([day, slot, activity, "Remplaçant", "Remplacement"])

    summary = workbook.create_sheet("Résumé")
    summary.append(["Médecin", "Créneaux affectés"])
    for name, variable in workload.items():
        summary.append([name, round(variable.value() or 0)])
    summary.append([])
    summary.append(["Statut", pulp.LpStatus[model.status]])
    summary.append(["Objectif", model.objective.value()])
    workbook.save(path)


def main() -> None:
    doctors, demand, absences = read_data(INPUT_FILE)
    model, assignment, replacement, workload = build_model(doctors, demand, absences)
    solver = pulp.PULP_CBC_CMD(msg=False)
    model.solve(solver)
    print(f"Statut: {pulp.LpStatus[model.status]}")
    print(f"Remplaçants: {sum(round(variable.value() or 0) for variable in replacement.values())}")
    for name, variable in workload.items():
        print(f"{name}: {round(variable.value() or 0)} créneaux")
    export_solution(OUTPUT_FILE, model, assignment, replacement, workload)
    print(f"Fichier créé: {OUTPUT_FILE.name}")


if __name__ == "__main__":
    main()